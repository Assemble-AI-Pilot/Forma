#!/usr/bin/env python3
"""Fill the Assemble CRC Consultant Recommendation template from a values JSON.

Why a script: the recommendation is the same template every time, so the fill is
deterministic. Doing it in code (rather than hand-editing XML each run) reliably
removes the yellow instruction prompts, keeps the table formatting intact, applies
the signature-threshold logic identically every time, and handles the list /
commitment / page-break formatting Assemble expects.

Usage:
    python fill_recommendation.py <values.json> <output.docx> [--template path]

The template defaults to ../assets/CRC_Recommendation_Template.docx (bundled).

CONVENTIONS
  * All dollar figures are EXCL GST (Assemble convention).
  * Only ONE of ABN / ACN is shown per entity (prefer ABN).
  * Remaining Budget is computed = budget - commitment total - other forecast.
  * Signature thresholds (ex GST): Sam Delmenico required when commitment total
    > $25,000; Carolyn Viney required when > $1,000,000. Jarvis is always Yes.
  * Any text wrapped in [square brackets] is treated as an unresolved/uncertain
    note and is HIGHLIGHTED YELLOW in the output so a reviewer spots it instantly.
  * "departures" and "attachments" render as a real numbered list.
  * The Recommendation for Signature table is forced onto its own page.
  * other_budget_forecast accepts "nil" (or "none", "na", "n/a", 0) and renders
    the standard nil entry, "$0 / NA".

VALUES JSON KEYS (all optional; an omitted field is left BLANK, not left as the
yellow prompt, so the reviewer sees exactly what still needs input -- the script
lists every blank cell when it finishes):
  project, discipline, agreement_type (usually one of AGREEMENT_TYPES; may be bespoke),
  assemble_entity, assemble_abn_acn, assemble_address,
  counterparty_entity, counterparty_abn_acn, counterparty_address,
  counterparty_signatory, counterparty_email, counterparty_mobile,
  scope, timelines, budget,
  cost_code_name, cost_code_number,
  commitment_stages  (list of [label, amount]),
  other_budget_forecast, comments, due_diligence,
  departures (list), public_liability, professional_indemnity,
  form_of_agreement, proposed_novated, design_consultant,
  attachments (list),
  sam_required / carolyn_required (optional overrides of the threshold logic)
"""
import argparse, json, os, re, sys

try:
    import docx
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx is required: pip install python-docx --break-system-packages")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(os.path.dirname(HERE), "assets", "CRC_Recommendation_Template.docx")
CELL_FONT_PT = 8

FIELD_CELLS = {
    "project": (0, 1), "discipline": (0, 2), "agreement_type": (0, 4),
    "assemble_entity": (1, 2), "assemble_abn_acn": (1, 3), "assemble_address": (1, 4),
    "counterparty_entity": (1, 7), "counterparty_abn_acn": (1, 8), "counterparty_address": (1, 9),
    "counterparty_signatory": (1, 10), "counterparty_email": (1, 11), "counterparty_mobile": (1, 12),
    "scope": (1, 14), "timelines": (1, 15),
    "budget": (1, 17), "cost_code_name": (1, 18), "cost_code_number": (1, 19),
    "other_budget_forecast": (1, 21), "remaining_budget": (1, 22), "comments": (1, 23),
    "due_diligence": (1, 25),
    "departures": (1, 27), "public_liability": (1, 28), "professional_indemnity": (1, 29),
    "form_of_agreement": (1, 30), "proposed_novated": (1, 31), "design_consultant": (1, 32),
    "attachments": (1, 34),
}
COMMITMENT_CELL = (1, 20)
SIGNATURE_CELLS = {"sam_required": (2, 3), "carolyn_required": (2, 4)}
LIST_FIELDS = {"departures", "attachments"}
# The three default agreements -- also the options of Forma's "Agreement Type" contract attribute.
# Anything else is a bespoke written agreement -- allowed, but it has no Forma option of its own.
AGREEMENT_TYPES = (
    "Major Design Consultant Agreement",
    "Minor Design Consultant Agreement",
    "Non-Design Consultant Agreement",
)
NIL_FORECAST = "$0 / NA"
NIL_WORDS = {"nil", "none", "na", "n/a", "0"}


def money(v):
    return f"${v:,.0f}" if isinstance(v, (int, float)) else str(v)

def parse_amount(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.sub(r"[^\d.]", "", str(v))
    return float(m) if m else None

def ensure_numbering(doc):
    """Create a fresh decimal numbering definition; return its numId."""
    numbering = doc.part.numbering_part.element
    abs_ids = [int(a.get(qn('w:abstractNumId'))) for a in numbering.findall(qn('w:abstractNum'))]
    new_abs = max(abs_ids) + 1 if abs_ids else 0
    a = OxmlElement('w:abstractNum'); a.set(qn('w:abstractNumId'), str(new_abs))
    lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), '0')
    for tag, val in (('w:start', '1'), ('w:numFmt', 'decimal'), ('w:lvlText', '%1.'), ('w:lvlJc', 'left')):
        e = OxmlElement(tag); e.set(qn('w:val'), val); lvl.append(e)
    pPr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '360'); ind.set(qn('w:hanging'), '360'); pPr.append(ind); lvl.append(pPr)
    a.append(lvl)
    first_num = numbering.find(qn('w:num'))
    numbering.insert(list(numbering).index(first_num), a) if first_num is not None else numbering.append(a)
    num_ids = [int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
    new_num = max(num_ids) + 1 if num_ids else 1
    n = OxmlElement('w:num'); n.set(qn('w:numId'), str(new_num))
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(new_abs)); n.append(ref)
    numbering.append(n)
    return new_num

def apply_number(p, num_id):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0'); numPr.append(ilvl)
    nid = OxmlElement('w:numId'); nid.set(qn('w:val'), str(num_id)); numPr.append(nid)
    pPr.insert(0, numPr)

def add_runs_with_highlight(p, text, bold=False):
    """Add text to paragraph; any [bracketed] span is highlighted yellow."""
    for i, seg in enumerate(re.split(r'(\[[^\]]*\])', str(text))):
        if seg == "":
            continue
        run = p.add_run(seg)
        run.font.size = Pt(CELL_FONT_PT)
        if bold:
            run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW if seg.startswith('[') else WD_COLOR_INDEX.AUTO

def _strip_para_mark_highlight(p):
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        return
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        return
    for h in rPr.findall(qn('w:highlight')):
        rPr.remove(h)

def reset_cell(cell):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    _strip_para_mark_highlight(p0)
    return p0

def blank_cell(cell):
    """Clear the template's yellow prompt and leave the cell empty."""
    p0 = reset_cell(cell)
    p0.style = "TableText"

def write_text(cell, value, doc=None, numbered=False):
    lines = value if isinstance(value, list) else [value]
    p0 = reset_cell(cell)
    num_id = ensure_numbering(doc) if numbered else None  # fresh list each time -> restarts at 1
    for i, line in enumerate(lines):
        p = p0 if i == 0 else cell.add_paragraph()
        p.style = "TableText"
        _strip_para_mark_highlight(p)
        if num_id is not None:
            apply_number(p, num_id)
        add_runs_with_highlight(p, line)

def write_commitment(cell, stages, total):
    p0 = reset_cell(cell)
    tab_pos = Inches(2.9)
    for i, (label, amount) in enumerate(stages):
        p = p0 if i == 0 else cell.add_paragraph()
        p.style = "TableText"
        p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.LEFT)
        r1 = p.add_run(f"{label}:"); r1.bold = True; r1.font.size = Pt(CELL_FONT_PT)
        r1.add_tab()
        r2 = p.add_run(money(amount)); r2.font.size = Pt(CELL_FONT_PT)
    p = cell.add_paragraph(); p.style = "TableText"
    p.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.LEFT)
    rt = p.add_run("Total"); rt.bold = True; rt.underline = True; rt.font.size = Pt(CELL_FONT_PT)
    rtab = p.add_run(); rtab.add_tab(); rtab.font.size = Pt(CELL_FONT_PT)   # tab NOT underlined
    rtot = p.add_run(money(total)); rtot.bold = True; rtot.underline = True; rtot.font.size = Pt(CELL_FONT_PT)

def set_required(cell, value):
    p0 = reset_cell(cell); p0.style = "TableText"
    run = p0.add_run(value); run.font.size = Pt(CELL_FONT_PT)
    run.font.highlight_color = WD_COLOR_INDEX.AUTO

def page_break_before_signature(doc):
    """Force table 2 (signatures) onto a new page."""
    sig_tbl = doc.tables[2]._tbl
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr'); r_pr = OxmlElement('w:rPr'); pPr.append(r_pr); p.append(pPr)
    r = OxmlElement('w:r'); br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); r.append(br); p.append(r)
    sig_tbl.addprevious(p)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("values"); ap.add_argument("output")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    args = ap.parse_args()

    v = json.load(open(args.values, encoding="utf-8-sig"))  # -sig: tolerate a Windows BOM
    at = v.get("agreement_type")
    if at and at not in AGREEMENT_TYPES:
        print(f"NOTE: agreement_type {at!r} is not one of the three defaults -- treated as a bespoke "
              f"agreement. Check the spelling if a default was meant; the Forma step needs an "
              f"Agreement Type option chosen.")
    ob = v.get("other_budget_forecast")
    if ob is not None and str(ob).strip().lower() in NIL_WORDS:
        v["other_budget_forecast"] = NIL_FORECAST

    d = docx.Document(args.template)
    tables = d.tables
    left_blank = []

    stages = v.get("commitment_stages"); total = None
    ci_t, ci_r = COMMITMENT_CELL
    if stages:
        total = sum(parse_amount(a) or 0 for _, a in stages)
        write_commitment(tables[ci_t].rows[ci_r].cells[1], stages, total)
    else:
        blank_cell(tables[ci_t].rows[ci_r].cells[1]); left_blank.append("commitment_stages")

    if "remaining_budget" not in v:
        b = parse_amount(v.get("budget")); other = parse_amount(v.get("other_budget_forecast")) or 0
        if b is not None and total is not None:
            v["remaining_budget"] = money(b - total - other)
    if "budget" in v:
        v["budget"] = money(v["budget"])

    for key, (ti, ri) in FIELD_CELLS.items():
        cell = tables[ti].rows[ri].cells[1]
        if key in v and v[key] not in (None, "", []):
            write_text(cell, v[key], doc=d, numbered=key in LIST_FIELDS)
        else:
            blank_cell(cell); left_blank.append(key)

    sam = v.get("sam_required"); carolyn = v.get("carolyn_required")
    if sam is None and total is not None:
        sam = "Yes" if total > 25000 else "No"
    if carolyn is None and total is not None:
        carolyn = "Yes" if total > 1000000 else "No"
    for key, value in (("sam_required", sam), ("carolyn_required", carolyn)):
        ti, ri = SIGNATURE_CELLS[key]
        cell = tables[ti].rows[ri].cells[2]
        if value is not None:
            set_required(cell, value)
        else:
            blank_cell(cell); left_blank.append(key)

    page_break_before_signature(d)
    d.save(args.output)
    print(f"Saved: {args.output}")
    if total is not None:
        print(f"Commitment total (ex GST): {money(total)} | Sam: {sam}  Carolyn: {carolyn}")
    if left_blank:
        print("Left blank (needs input): " + ", ".join(left_blank))
    else:
        print("Every field filled.")

if __name__ == "__main__":
    main()
