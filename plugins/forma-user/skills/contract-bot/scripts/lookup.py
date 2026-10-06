#!/usr/bin/env python3
"""Resolve Assemble developer entities and consultant cost codes for Gem Bot.

Reads directly from the source spreadsheets so results stay current with whatever
workbook is bundled (or whatever newer copy the user supplies via --bible / --codes).

Usage:
    python lookup.py entity "<project search term>"      # e.g. "Carinish" or "Footscray"
    python lookup.py entity --list                        # list every Project-Developer entity
    python lookup.py costcode "<discipline search term>"  # e.g. "structural" or "town planner"
    python lookup.py costcode --list                      # list every cost code

The bundled Entity_Bible.xlsx is an EXTRACT of the master (make_entity_extract.py):
only the columns read here, with no tax file numbers, corporate keys or officeholders.

Optional overrides:
    --bible /path/to/Entity_Bible.xlsx
    --codes /path/to/Consultant_Cost_Codes.xlsx

Design notes:
  * Entity matching filters to rows where Purpose contains "Project - Developer"
    (the only entities valid as the Assemble counterparty on a consultant agreement).
  * A project can map to MORE THAN ONE developer entity. The script prints ALL
    matches; Gem Bot must ask the user to pick rather than guessing.
  * Fields are cleaned (collapsed whitespace, stripped). A blank Registered Address
    is reported as "(blank in bible - confirm with user)" so it is never silently dropped.
"""
import argparse
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl --break-system-packages")

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
DEFAULT_BIBLE = os.path.join(ASSETS, "Entity_Bible.xlsx")
DEFAULT_CODES = os.path.join(ASSETS, "Consultant_Cost_Codes.xlsx")

# Column map for the "Summary of Assemble Entities" sheet (header row 6, data from row 7)
ENTITY_SHEET = "Summary of Assemble Entities"
COL = {"name": 1, "purpose": 4, "project_address": 5, "acn": 6, "abn": 7, "registered_address": 11}
HEADER_ROW = 6


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def load_developer_entities(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[ENTITY_SHEET]
    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        purpose = clean(ws.cell(row=r, column=COL["purpose"]).value)
        if "Project - Developer" not in purpose:
            continue
        rows.append({
            "name": clean(ws.cell(row=r, column=COL["name"]).value),
            "project_address": clean(ws.cell(row=r, column=COL["project_address"]).value),
            "acn": clean(ws.cell(row=r, column=COL["acn"]).value),
            "abn": clean(ws.cell(row=r, column=COL["abn"]).value),
            "registered_address": clean(ws.cell(row=r, column=COL["registered_address"]).value),
        })
    return rows


def print_entity(e):
    print(f"  Assemble Entity      : {e['name']}")
    print(f"  Project Address      : {e['project_address'] or '(none recorded)'}")
    print(f"  ACN                  : {e['acn'] or '(none recorded)'}")
    print(f"  ABN                  : {e['abn'] or '(none recorded)'}")
    addr = e["registered_address"] or "(blank in bible - confirm with user)"
    print(f"  Registered Address   : {addr}")
    print()


def cmd_entity(args):
    rows = load_developer_entities(args.bible)
    if args.list or not args.term:
        print(f"All {len(rows)} Project-Developer entities:\n")
        for e in sorted(rows, key=lambda x: x["name"]):
            print_entity(e)
        return
    term = args.term.lower()
    matches = [e for e in rows
               if term in e["name"].lower() or term in e["project_address"].lower()]
    if not matches:
        print(f'No Project-Developer entity matched "{args.term}".')
        print("Try a broader term, or run:  python lookup.py entity --list")
        return
    if len(matches) > 1:
        print(f'{len(matches)} entities matched "{args.term}" - ASK THE USER WHICH ONE:\n')
    else:
        print(f'1 entity matched "{args.term}":\n')
    for e in matches:
        print_entity(e)


def load_cost_codes(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Cost Codes"]
    codes = []
    for r in range(1, ws.max_row + 1):
        code = clean(ws.cell(row=r, column=1).value)
        name = clean(ws.cell(row=r, column=2).value)
        # valid code rows look like 17141.03101
        if re.match(r"^\d{5}\.\d{4,5}$", code):
            codes.append({"code": code, "name": name})
    return codes


def cmd_costcode(args):
    codes = load_cost_codes(args.codes)
    if args.list or not args.term:
        print(f"All {len(codes)} cost codes:\n")
        for c in codes:
            print(f"  {c['code']}  {c['name']}")
        return
    term = args.term.lower()
    matches = [c for c in codes if term in c["name"].lower()]
    if not matches:
        print(f'No cost code matched "{args.term}". Run:  python lookup.py costcode --list')
        return
    if len(matches) > 1:
        print(f'{len(matches)} cost codes matched "{args.term}" - confirm the right discipline:\n')
    else:
        print(f'1 cost code matched "{args.term}":\n')
    for c in matches:
        print(f"  Code (Number) : {c['code']}")
        print(f"  Code (Name)   : {c['name']}")
        print()


def main():
    p = argparse.ArgumentParser(description="Gem Bot entity / cost-code lookup")
    sub = p.add_subparsers(dest="mode", required=True)

    # `--list` is a flag, not a value for `term`: argparse reads anything starting
    # with "-" as an option, so `entity --list` used to fail with "term required".
    pe = sub.add_parser("entity")
    pe.add_argument("term", nargs="?", help="project search term")
    pe.add_argument("--list", action="store_true", help="list every Project-Developer entity")
    pe.add_argument("--bible", default=DEFAULT_BIBLE)
    pe.set_defaults(func=cmd_entity)

    pc = sub.add_parser("costcode")
    pc.add_argument("term", nargs="?", help="discipline search term")
    pc.add_argument("--list", action="store_true", help="list every cost code")
    pc.add_argument("--codes", default=DEFAULT_CODES)
    pc.set_defaults(func=cmd_costcode)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
